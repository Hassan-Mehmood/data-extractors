from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


_HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
_HEADER_FONT = Font(bold=True, color="FFFFFF")


def _write_table_sheet(wb: Workbook, df: pd.DataFrame, sheet_name: str) -> None:
    ws = wb.create_sheet(title=sheet_name)

    # Write header row
    for col_idx, col_name in enumerate(df.columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL

    # Write data rows
    for row_idx, row in enumerate(df.itertuples(index=False), start=2):
        for col_idx, value in enumerate(row, start=1):
            ws.cell(row=row_idx, column=col_idx, value=value)

    # Freeze header row
    ws.freeze_panes = "A2"

    # Auto-fit column widths (capped at 60)
    for col_idx, col_name in enumerate(df.columns, start=1):
        max_len = max(
            len(str(col_name)),
            *(len(str(v)) for v in df.iloc[:, col_idx - 1]),
            0,
        )
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 60)


def _write_overview_sheet(wb: Workbook, tables: list[pd.DataFrame]) -> None:
    ws = wb.create_sheet(title="Overview", index=0)

    headers = ["#", "Sheet", "Rows", "Columns"]
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL

    ws.freeze_panes = "A2"

    for i, df in enumerate(tables, start=1):
        ws.append([i, f"Table {i}", len(df), len(df.columns)])

    for col_idx in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 16


def export_to_excel(tables: list[pd.DataFrame], output_path: Path) -> None:
    """
    Write a list of DataFrames to an Excel workbook at output_path.

    Sheet layout:
      - Sheet 0: "Overview" — index of all tables with row/column counts
      - Sheets 1‥N: "Table 1", "Table 2", … — one sheet per table
    """
    wb = Workbook()
    # Remove the default empty sheet created by openpyxl
    wb.remove(wb.active)

    _write_overview_sheet(wb, tables)
    for i, df in enumerate(tables, start=1):
        _write_table_sheet(wb, df, f"Table {i}")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(output_path))
